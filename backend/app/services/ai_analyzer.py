"""AI解析パイプライン

アップロードされた書類（PDF/画像/Excel/Word）をClaudeで解析し、
構造化データ＋信頼度スコアを ai_analyses に保存する。

プロバイダは3種類:
- anthropic : Anthropic API（ANTHROPIC_API_KEY 必須。本番用）
- claude_cli: ローカルの claude コマンド（開発機のClaude Code認証を利用）
- mock      : 決定的なダミー解析（APIなしでの動作確認・テスト用）
"""
import asyncio
import base64
import json
import logging
import os
import re
import shutil
import time
import uuid

from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.core.config import settings
from app.core.database import async_session
from app.models.ai_analysis import AIAnalysis
from app.models.document import Document
from app.models.notification import Notification
from app.models.user import User

logger = logging.getLogger("ai_analyzer")

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".tiff", ".heic", ".heif"}
IMAGE_MEDIA_TYPES = {
    ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
    ".tiff": "image/tiff", ".heic": "image/heic", ".heif": "image/heif",
}

ANALYSIS_PROMPT = """あなたは展示会の事務局スタッフとして、出展社から提出された書類を読み取り、構造化データに変換します。

## 書類の情報
- 提出カテゴリ: {category_name}（{category_description}）
- 元ファイル名: {file_name}

## 出力形式
以下のJSONだけを出力してください。説明文・コードフェンスは不要です。

{{
  "document_type": "order | design | contract | other",
  "detected_company_name": "書類から読み取れた会社名（なければnull）",
  "detected_booth_number": "書類から読み取れたブース番号（なければnull）",
  "summary": "書類の内容の1〜2文の要約（日本語）",
  "order_items": [
    {{"item_name": "品目名", "quantity": 数値, "unit": "単位", "unit_price": 数値orNull, "total_price": 数値orNull}}
  ],
  "total_amount": 合計金額の数値（なければnull）,
  "delivery_date": "YYYY-MM-DD形式の希望納期（なければnull）",
  "special_instructions": "特記事項（なければnull）",
  "extracted_text": "書類全体のテキスト書き起こし（最大2000文字）",
  "field_confidence": {{
    "order_items": 0.0〜1.0,
    "total_amount": 0.0〜1.0,
    "detected_company_name": 0.0〜1.0,
    "delivery_date": 0.0〜1.0
  }},
  "overall_confidence": 0.0〜1.0,
  "extraction_notes": "読み取り時の注意点・曖昧だった箇所（日本語）"
}}

## ルール
- 注文書・申込書でなければ order_items は空配列にする
- 読み取れない値は推測せず null にして、該当フィールドの confidence を下げる
- 手書き・低解像度・傾きなどで判読が怪しい場合は overall_confidence を正直に下げる
- 金額は数値のみ（円記号・カンマ除去）
"""


def extract_office_text(file_path: str) -> str | None:
    """Excel/Wordからテキストを抽出する。対応外の形式はNone"""
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"### シート: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        lines.append("\t".join(cells))
            return "\n".join(lines)[:20000]
        if ext in (".docx", ".doc"):
            import docx
            d = docx.Document(file_path)
            lines = [p.text for p in d.paragraphs if p.text.strip()]
            for table in d.tables:
                for row in table.rows:
                    lines.append("\t".join(c.text for c in row.cells))
            return "\n".join(lines)[:20000]
    except Exception as e:
        logger.warning("office text extraction failed: %s", e)
    return None


def parse_json_response(text: str) -> dict:
    """LLM応答からJSONを頑健に取り出す"""
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    start = text.find("{")
    end = text.rfind("}")
    if start == -1 or end == -1:
        raise ValueError(f"JSONが見つかりません: {text[:200]}")
    return json.loads(text[start:end + 1])


def resolve_provider() -> str:
    if settings.ai_provider != "auto":
        return settings.ai_provider
    if settings.anthropic_api_key or os.environ.get("ANTHROPIC_API_KEY"):
        return "anthropic"
    if shutil.which("claude"):
        return "claude_cli"
    return "mock"


async def run_anthropic(prompt: str, file_path: str, office_text: str | None) -> tuple[dict, dict]:
    """Anthropic APIで解析。(結果JSON, メタ情報)を返す"""
    import anthropic

    client = anthropic.AsyncAnthropic(
        api_key=settings.anthropic_api_key or os.environ.get("ANTHROPIC_API_KEY")
    )
    ext = os.path.splitext(file_path)[1].lower()
    content: list = []
    if office_text is not None:
        content.append({"type": "text", "text": f"## 書類テキスト（{ext}から抽出）\n{office_text}"})
    elif ext == ".pdf":
        with open(file_path, "rb") as f:
            data = base64.standard_b64encode(f.read()).decode()
        content.append({"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": data}})
    elif ext in IMAGE_EXTS:
        with open(file_path, "rb") as f:
            data = base64.standard_b64encode(f.read()).decode()
        content.append({"type": "image", "source": {"type": "base64", "media_type": IMAGE_MEDIA_TYPES[ext], "data": data}})
    else:
        raise ValueError(f"解析対象外の形式です: {ext}")
    content.append({"type": "text", "text": prompt})

    resp = await client.messages.create(
        model=settings.ai_model,
        max_tokens=4096,
        messages=[{"role": "user", "content": content}],
    )
    result = parse_json_response(resp.content[0].text)
    meta = {
        "llm_model": settings.ai_model,
        "llm_prompt_tokens": resp.usage.input_tokens,
        "llm_completion_tokens": resp.usage.output_tokens,
    }
    return result, meta


async def run_claude_cli(prompt: str, file_path: str, office_text: str | None) -> tuple[dict, dict]:
    """ローカルのclaudeコマンドで解析（開発用）"""
    if office_text is not None:
        full_prompt = f"## 書類テキスト（Office形式から抽出）\n{office_text}\n\n{prompt}"
        cmd = ["claude", "-p", full_prompt]
    else:
        full_prompt = f"次のファイルを読み取って解析してください: {os.path.abspath(file_path)}\n\n{prompt}"
        cmd = ["claude", "-p", full_prompt, "--allowedTools", "Read"]

    proc = await asyncio.create_subprocess_exec(
        *cmd,
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.PIPE,
    )
    try:
        stdout, stderr = await asyncio.wait_for(proc.communicate(), timeout=300)
    except asyncio.TimeoutError:
        proc.kill()
        raise RuntimeError("claude CLIがタイムアウトしました")
    if proc.returncode != 0:
        raise RuntimeError(f"claude CLIが失敗しました: {stderr.decode()[:500]}")
    result = parse_json_response(stdout.decode())
    return result, {"llm_model": "claude-cli"}


def run_mock(category_name: str, file_name: str) -> tuple[dict, dict]:
    """決定的なダミー解析（APIなしでフロー全体を確認するため）"""
    # ファイル名に「手書き」「低画質」を含む場合は低信頼度にしてレビューフローを通す
    low_quality = any(k in file_name for k in ("手書き", "低画質", "review"))
    conf = 0.62 if low_quality else 0.93
    is_order = any(k in category_name for k in ("申込", "注文", "レンタル"))
    result = {
        "document_type": "order" if is_order else "design",
        "detected_company_name": "出展社A株式会社",
        "detected_booth_number": "A-01",
        "summary": f"{category_name}の提出書類（モック解析）。",
        "order_items": [
            {"item_name": "パネル（白） W900×H2400", "quantity": 10, "unit": "枚", "unit_price": 5000, "total_price": 50000},
            {"item_name": "スポットライト 100W", "quantity": 4, "unit": "台", "unit_price": 3000, "total_price": 12000},
        ] if is_order else [],
        "total_amount": 62000 if is_order else None,
        "delivery_date": "2026-08-31" if is_order else None,
        "special_instructions": "搬入は前日16時以降希望" if is_order else None,
        "extracted_text": f"（モック）{file_name} の解析テキスト",
        "field_confidence": {
            "order_items": conf, "total_amount": conf,
            "detected_company_name": min(conf + 0.05, 1.0), "delivery_date": conf,
        },
        "overall_confidence": conf,
        "extraction_notes": "モックプロバイダによる解析です" + ("（低品質ファイルを検出）" if low_quality else ""),
    }
    return result, {"llm_model": "mock"}


async def analyze_document(document_id: uuid.UUID) -> None:
    """バックグラウンドで1ドキュメントを解析する（独自DBセッションを持つ）"""
    async with async_session() as db:
        result = await db.execute(
            select(Document)
            .options(selectinload(Document.submission_category))
            .where(Document.id == document_id)
        )
        document = result.scalar_one_or_none()
        if document is None:
            logger.error("document not found: %s", document_id)
            return

        document.status = "processing"
        await db.commit()

        category = document.submission_category
        prompt = ANALYSIS_PROMPT.format(
            category_name=category.name if category else "不明",
            category_description=(category.description or "") if category else "",
            file_name=document.file_name,
        )

        provider = resolve_provider()
        started = time.monotonic()
        try:
            office_text = extract_office_text(document.storage_path)
            if provider == "anthropic":
                data, meta = await run_anthropic(prompt, document.storage_path, office_text)
            elif provider == "claude_cli":
                data, meta = await run_claude_cli(prompt, document.storage_path, office_text)
            else:
                data, meta = run_mock(category.name if category else "", document.file_name)
        except Exception:
            logger.exception("AI解析に失敗しました: %s", document_id)
            document.status = "error"
            await db.commit()
            return

        elapsed_ms = int((time.monotonic() - started) * 1000)

        field_conf: dict = data.get("field_confidence") or {}
        confidence = data.get("overall_confidence")
        if confidence is None:
            confidence = min(field_conf.values(), default=0.5)
        confidence = max(0.0, min(1.0, float(confidence)))
        low_fields = [k for k, v in field_conf.items() if isinstance(v, (int, float)) and v < 0.8]

        auto_ok = confidence >= settings.auto_approve_threshold
        analysis = AIAnalysis(
            document_id=document.id,
            extracted_text=data.get("extracted_text"),
            structured_data=data,
            confidence_score=round(confidence, 2),
            low_confidence_fields=low_fields,
            processing_time_ms=elapsed_ms,
            review_status="auto_approved" if auto_ok else "pending_review",
            llm_model=meta.get("llm_model"),
            llm_prompt_tokens=meta.get("llm_prompt_tokens"),
            llm_completion_tokens=meta.get("llm_completion_tokens"),
        )
        db.add(analysis)

        document.document_category = data.get("document_type") or "other"
        document.status = "analyzed" if auto_ok else "review_needed"
        await db.commit()

        # 高信頼度の注文書はレビューを待たず注文データ化する
        if auto_ok:
            from app.services.order_builder import create_order_from_analysis
            await create_order_from_analysis(db, document, analysis)
            await db.commit()

        await _notify_recipients(db, document, analysis)


async def _notify_recipients(db, document: Document, analysis: AIAnalysis) -> None:
    """受取先組織＋主催者のユーザーに通知を作成する"""
    from app.models.organization import Organization

    org_ids = {document.recipient_org_id}
    result = await db.execute(select(Organization.id).where(Organization.org_type == "organizer"))
    org_ids.update(result.scalars().all())

    users = (await db.execute(select(User).where(User.organization_id.in_(org_ids), User.is_active.is_(True)))).scalars().all()
    needs_review = analysis.review_status == "pending_review"
    for u in users:
        db.add(Notification(
            user_id=u.id,
            type="review_needed" if needs_review else "document_received",
            title=("要レビュー: " if needs_review else "新着書類: ") + document.file_name,
            message=f"信頼度 {float(analysis.confidence_score):.0%} で解析されました。",
            reference_type="document",
            reference_id=document.id,
        ))
    await db.commit()
